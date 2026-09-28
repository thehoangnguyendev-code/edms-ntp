from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.section import WD_SECTION
from pathlib import Path

OUT = Path(r"D:\edms-project\docs\deliverables\EQMS_Notification_Inventory_and_Gap_Analysis.docx")

BLUE = "2E74B5"
DARK = "1F4D78"
INK = "1F2937"
MUTED = "5B6472"
HEADER_FILL = "E8EEF5"
LIGHT_FILL = "F4F6F9"
GREEN = "1F6B45"
AMBER = "7A5A00"
RED = "9B1C1C"

def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)

def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    mar = tc_pr.first_child_found_in("w:tcMar")
    if mar is None:
        mar = OxmlElement("w:tcMar")
        tc_pr.append(mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")

def set_cell_width(cell, width):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width))
    tc_w.set(qn("w:type"), "dxa")

def set_table_geometry(table, widths, indent=120):
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent))
    tbl_ind.set(qn("w:type"), "dxa")
    layout = tbl_pr.first_child_found_in("w:tblLayout")
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            set_cell_width(cell, widths[idx])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    tr_pr.append(el)

def font(run, size=10, bold=False, color=INK, italic=False):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)

def style_doc(doc):
    sec = doc.sections[0]
    sec.top_margin = Inches(1)
    sec.bottom_margin = Inches(1)
    sec.left_margin = Inches(1)
    sec.right_margin = Inches(1)
    sec.header_distance = Inches(0.492)
    sec.footer_distance = Inches(0.492)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15
    for name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK, 8, 4),
    ]:
        st = doc.styles[name]
        st.font.name = "Calibri"
        st._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        st._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        st.font.size = Pt(size)
        st.font.color.rgb = RGBColor.from_string(color)
        st.font.bold = True
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)

def p(doc, text="", bold_prefix=None, color=INK, size=10.5, italic=False):
    para = doc.add_paragraph()
    if bold_prefix and text.startswith(bold_prefix):
        r = para.add_run(bold_prefix)
        font(r, size, True, color)
        r = para.add_run(text[len(bold_prefix):])
        font(r, size, False, color, italic)
    else:
        r = para.add_run(text)
        font(r, size, False, color, italic)
    return para

def bullet(doc, text):
    para = doc.add_paragraph(style="List Bullet")
    para.paragraph_format.space_after = Pt(4)
    r = para.add_run(text)
    font(r, 10.5)
    return para

def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.style = "Table Grid"
    set_table_geometry(table, widths)
    header = table.rows[0]
    set_repeat_table_header(header)
    for i, header_text in enumerate(headers):
        cell = header.cells[i]
        set_cell_shading(cell, HEADER_FILL)
        para = cell.paragraphs[0]
        para.paragraph_format.space_after = Pt(0)
        run = para.add_run(header_text)
        font(run, 9, True, DARK)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cell = cells[i]
            para = cell.paragraphs[0]
            para.paragraph_format.space_after = Pt(0)
            # Status text gets intentional color.
            col = INK
            if str(value).startswith("Co san"):
                col = GREEN
            elif str(value).startswith("Cau hinh"):
                col = AMBER
            elif str(value).startswith("Chua co"):
                col = RED
            run = para.add_run(str(value))
            font(run, 8.4, i == 0 and len(str(value)) < 42, col)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table

def lead_box(doc, text):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360])
    cell = table.cell(0,0)
    set_cell_shading(cell, LIGHT_FILL)
    para = cell.paragraphs[0]
    para.paragraph_format.space_after = Pt(0)
    r = para.add_run(text)
    font(r, 10.5, False, DARK)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

def footer(section):
    para = section.footer.paragraphs[0]
    para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = para.add_run("EQMS - Notification inventory | Internal working document")
    font(r, 8, False, MUTED)

def main():
    doc = Document()
    style_doc(doc)
    footer(doc.sections[0])

    title = doc.add_paragraph()
    title.paragraph_format.space_before = Pt(10)
    title.paragraph_format.space_after = Pt(4)
    r = title.add_run("EQMS NOTIFICATION INVENTORY")
    font(r, 23, True, DARK)
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(14)
    r = subtitle.add_run("Danh muc thong bao hien co, cau hinh chua co trigger, va de xuat bo sung")
    font(r, 12, False, MUTED)
    for label, value in [("Pham vi", "Backend notification/email flows va frontend inbox; danh gia theo source hien tai"),
                         ("Muc dich", "Lam co so de viet yeu cau nghiep vu va backlog notification"),
                         ("Ngay lap", "12/08/2026")]:
        pp = doc.add_paragraph()
        pp.paragraph_format.space_after = Pt(2)
        a = pp.add_run(label + ": ")
        font(a, 10, True, DARK)
        b = pp.add_run(value)
        font(b, 10, False, INK)
    lead_box(doc, "Ket luan nhanh: he thong da co pipeline inbox/realtime va email theo template/policy. Mot so event da duoc catalog hoa nhung can xac nhan trigger runtime; danh sach de xuat o Phan 4 uu tien cac diem cham workflow, deadline va compliance.")

    doc.add_heading("1. Cach doc tai lieu", level=1)
    p(doc, "Trang thai trong bang duoc phan loai nhu sau:")
    for text in [
        "Co san - trigger trong source va co kenh delivery (in-app, email, hoac ca hai).",
        "Cau hinh - event/policy/template da co, nhung trong pham vi ra soat chua xac nhan trigger nghiep vu thuc te.",
        "Chua co - nhu cau de xuat, can bo sung event, template, recipient rule va trigger.",
        "Thong bao he thong/security ap dung cho tai khoan, ha tang hoac cau hinh; thong bao nghiep vu gan voi workflow module.",
    ]:
        bullet(doc, text)
    p(doc, "Luu y: mot workflow email co the dong thoi tao inbox notification. Neu workflow da qua Notification Policy, he thong co co che tranh tao inbox trung lap.")

    doc.add_heading("2. Thong bao dang co trigger", level=1)
    existing = [
        ("Document Control", "Submit for review", "Nghiep vu", "In-app; Email neu policy EMAIL duoc bat", "Reviewer"),
        ("Document Control", "Review completed", "Nghiep vu", "In-app; Email neu policy bat", "Author, Approver"),
        ("Document Control", "Approved", "Nghiep vu", "In-app; Email neu policy bat", "Author, Owner"),
        ("Document Control", "Published / Effective", "GMP mandatory", "In-app + email workflow", "Author, Reviewer, Approver / stakeholder"),
        ("Document Control", "Author upload to Office Online thanh cong", "Nghiep vu", "Email + inbox", "Author va toan bo Co-author cua revision"),
        ("Document Control", "Author Complete Editing thanh cong", "Nghiep vu", "Email + inbox", "DCO / Document Controller co quyen Submit for Review"),
        ("Document Control", "Revision cancelled", "Nghiep vu", "Email + inbox", "Author va Co-author; co ly do cancel"),
        ("Document Control", "Reviewer / Approver ke tiep can xu ly", "Nghiep vu", "Email + inbox", "Reviewer hoac Approver ke tiep"),
        ("Document Control", "Review rejected", "Nghiep vu", "Email + inbox", "Workflow Coordinator"),
        ("Document Control", "Ready for publishing", "Nghiep vu", "Email + inbox", "Workflow Coordinator / DCO"),
        ("Document Control", "Training required sau approval", "Nghiep vu", "Email + inbox", "Revision stakeholders"),
        ("Controlled Copies", "Copy distributed", "GMP mandatory", "In-app + email", "Recipient / holder"),
        ("Controlled Copies", "Copy recalled", "GMP mandatory", "In-app + email", "Recipient / holder"),
        ("Controlled Copies", "Copy workflow update: cancel / recall / ...", "Nghiep vu", "Email + inbox", "Stakeholder da tham gia thao tac"),
        ("Controlled Copies", "Expiry reminder (7 ngay truoc han)", "Nghiep vu", "Email; event in-app da catalog", "Holder va DCO co quyen quan ly workspace"),
        ("Security", "Password changed / reset by admin", "System / GMP", "In-app; reset link la email", "Chu tai khoan"),
        ("Security", "MFA enabled / disabled", "System / GMP", "In-app", "Chu tai khoan"),
        ("Security", "Account locked", "System / GMP", "In-app", "Chu tai khoan"),
        ("Security", "Account suspended / terminated", "System / security", "In-app; email neu policy bat", "Nguoi dung bi suspend / terminate"),
        ("Security", "Access review campaign due", "Security / compliance", "In-app; email neu policy bat", "Reviewer cua campaign"),
        ("User Management", "Microsoft invitation accepted", "Nghiep vu", "In-app", "Admin co quyen xem external provisioning"),
        ("Authentication", "Password reset link", "System / security", "Email only", "Nguoi yeu cau reset"),
        ("Authentication", "MFA OTP / login code", "System / security", "Email only", "Nguoi dang xac thuc"),
    ]
    add_table(doc, ["Module", "Su kien", "Loai", "Kenh", "Nguoi nhan"], existing, [1600, 2450, 1450, 1700, 2160])

    doc.add_heading("3. Da cau hinh/catalog nhung can xac nhan trigger", level=1)
    configured = [
        ("Cau hinh - can xac nhan", "Document Control", "Periodic review due", "In-app", "Owner, affected users", "Event catalog co; can xac nhan job/trigger runtime."),
        ("Cau hinh - can xac nhan", "Audit Trail", "Audit review campaign assigned", "In-app", "Assignee", "Event catalog co; can xac nhan call site."),
        ("Cau hinh - can xac nhan", "Audit Trail", "Audit export completed", "In-app", "Nguoi yeu cau export / owner", "Event catalog co; can xac nhan call site."),
        ("Cau hinh - can xac nhan", "System", "Maintenance scheduled", "In-app", "Tat ca Active users", "Event catalog co; can xac nhan man hinh/job phat thong bao."),
        ("Cau hinh - can xac nhan", "Controlled Copies", "Copy destroyed", "In-app", "Recipient va Owner", "Event catalog co; can xac nhan call site."),
    ]
    add_table(doc, ["Trang thai", "Module", "Thong bao", "Kenh", "Nguoi nhan", "Ghi chu"], configured, [1350, 1200, 1900, 1100, 1750, 2060])

    doc.add_heading("4. Thong bao chua co - de xuat backlog", level=1)
    proposals = [
        ("Chua co - P1", "Document Control", "Author / Co-author duoc chi dinh hoac thay doi", "Nghiep vu", "In-app + email", "Nguoi moi duoc them/doi vai tro; DCO khi thay doi do nguoi khac thuc hien"),
        ("Chua co - P1", "Document Control", "Reviewer / Approver sap qua han", "Nghiep vu", "In-app + email", "Nguoi dang cho xu ly; DCO khi con 1 ngay"),
        ("Chua co - P1", "Document Control", "Reviewer / Approver qua han va escalation", "Compliance", "In-app + email", "Nguoi qua han, DCO, Workflow Coordinator, Author"),
        ("Chua co - P1", "Training", "Training sap han / qua han", "GMP / compliance", "In-app + email", "Trainee; manager va Training Coordinator khi qua han"),
        ("Chua co - P1", "Controlled Copies", "Chua acknowledge trong SLA", "GMP / compliance", "In-app + email", "Holder; DCO khi escalation"),
        ("Chua co - P2", "Document Control", "Document sap het hieu luc / obsolete", "Nghiep vu", "In-app + email", "Owner, Author, DCO, holder copy lien quan"),
        ("Chua co - P2", "Document Control", "Reviewer/Approver assignment da thay doi", "Audit-sensitive", "In-app + email", "Nguoi bi them/xoa, Author, DCO"),
        ("Chua co - P2", "Controlled Copies", "Distribution batch thanh cong / that bai mot phan", "Nghiep vu", "In-app + email", "DCO, distributor, requester"),
        ("Chua co - P2", "Audit Trail", "Audit review campaign sap/da qua han", "Compliance", "In-app + email", "Assignee; QA/Compliance owner khi escalation"),
        ("Chua co - P2", "Security", "Dang nhap bat thuong", "System / security", "In-app + email", "Chu tai khoan; Security Admin neu risk cao"),
        ("Chua co - P2", "System", "Email delivery / scheduler failure", "System", "In-app + email / monitoring", "System Admin, module owner"),
        ("Chua co - P3", "System", "Storage/DB sap day; backup that bai", "System", "In-app + email / monitoring", "System Admin"),
    ]
    add_table(doc, ["Uu tien", "Module", "Thong bao", "Loai", "Kenh", "Nguoi nhan"], proposals, [1150, 1200, 2050, 1250, 1350, 2360])

    doc.add_heading("5. Chi tiet 3 thong bao can chap thuan yeu cau", level=1)
    details = [
        ("A. Chi dinh Author / Co-author", "Trigger: tao document/revision hoac DCO luu thay doi Author/Co-author. Chi gui cho nguoi vua duoc them/doi vai tro; khong gui lai cho user khong thay doi. Noi dung: document number, title, revision, vai tro moi, actor, action URL. Kenh: in-app + email. Loai: nghiep vu ca nhan."),
        ("B. Upload to Office Online thanh cong", "Trigger: Author upload working copy thanh cong. Recipient: tat ca Co-author (co the gui ca Author nhu confirmation, nhung khong bat buoc). Noi dung: document/revision, ai upload, office edit URL, thoi gian. Kenh hien tai: email + inbox. Loai: nghiep vu ca nhan."),
        ("C. Complete Editing thanh cong", "Trigger: Author complete editing, file da sync/lock va Prepared signature da ghi. Recipient: DCO/Document Controller co quyen SUBMIT_FOR_REVIEW theo workflow policy cua document type. Noi dung: document/revision, Author, trang thai Draft ready for DCO check, action URL. Kenh hien tai: email + inbox. Loai: nghiep vu ca nhan."),
    ]
    add_table(doc, ["Thong bao", "Yeu cau / quy tac"], details, [2350, 7010])

    doc.add_heading("6. Quy tac chung de viet yeu cau tiep theo", level=1)
    for text in [
        "Moi notification can co event code, trigger, recipient rule, kenh, priority, title/body template, action URL va data variables.",
        "Phan biet recipient theo vai tro cua revision hien tai va document master; uu tien snapshot workflow cua revision neu su kien gan voi revision.",
        "Tranh spam: deduplicate theo event + entity + recipient; chi gui assignment notification khi co thay doi thuc te.",
        "Voi GMP/security: ghi audit trail, luu delivery failure, va khong cho user tat neu event mandatory.",
        "Escalation can co SLA, moc nhac (vi du T-3, T-1, overdue) va audience mo rong ro rang.",
    ]:
        bullet(doc, text)

    doc.add_heading("7. Ghi chu ve cau hinh delivery", level=1)
    p(doc, "Notification Policy quyet dinh event co active khong, enabled channels va recipient rules. User preference chi ap dung cho event khong mandatory. Email can co active template/version, recipient co email hop le va email notification duoc bat. Inbox notification duoc luu trong Notifications va day realtime sau khi transaction thanh cong.")

    doc.save(OUT)
    print(OUT)

if __name__ == "__main__":
    main()
